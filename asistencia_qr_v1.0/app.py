#VERSION = "1.0"
import tkinter as tk

from pathlib import Path
from datetime import datetime

from PIL import (
    Image,
    ImageTk,
    ImageEnhance,
    ImageOps,
    ImageDraw,
    ImageFont
)

from openpyxl import Workbook, load_workbook


# ==========================================
# CONFIGURACIÓN
# ==========================================

BASE_DIR = Path(__file__).resolve().parent

ALUMNOS_FILE = BASE_DIR / "alumnos.xlsx"
ASISTENCIAS_DIR = BASE_DIR / "asistencias"

ASSETS_DIR = BASE_DIR / "assets"

FONDO_FILE = ASSETS_DIR / "fondo.png"
LOGO_FILE = ASSETS_DIR / "logo.png"


# ==========================================
# COLORES
# ==========================================

COLOR_PANEL = "#0a172c"
COLOR_INPUT = "#101f36"
COLOR_AZUL = "#3d8bfd"

COLOR_TEXTO = "#ffffff"
COLOR_SECUNDARIO = "#aebbd0"

COLOR_EXITO = "#70d6a3"
COLOR_ERROR = "#ff6b6b"
COLOR_ADVERTENCIA = "#ffb86c"


# ==========================================
# PREPARAR CARPETAS
# ==========================================

ASISTENCIAS_DIR.mkdir(exist_ok=True)


# ==========================================
# CREAR ARCHIVO DE ALUMNOS
# ==========================================

def crear_archivo_alumnos():

    if ALUMNOS_FILE.exists():
        return

    workbook = Workbook()

    hoja = workbook.active
    hoja.title = "Alumnos"

    hoja.append([
        "Matrícula",
        "Nombre",
        "Carrera",
        "Nivel"
    ])

    workbook.save(ALUMNOS_FILE)
    workbook.close()


# ==========================================
# OBTENER ARCHIVO DE ASISTENCIA DEL DÍA
# ==========================================

def obtener_archivo_asistencia():

    fecha_actual = datetime.now().strftime("%Y-%m-%d")

    return ASISTENCIAS_DIR / f"asistencia_{fecha_actual}.xlsx"


# ==========================================
# GENERAR REGISTRO DEL DÍA
# ==========================================

def generar_archivo_dia():

    archivo = obtener_archivo_asistencia()

    if archivo.exists():

        mensaje.config(
            text="El registro de hoy ya existe.",
            fg=COLOR_SECUNDARIO
        )

    else:

        workbook = Workbook()

        hoja = workbook.active
        hoja.title = "Asistencias"

        hoja.append([
            "Matrícula",
            "Nombre",
            "Carrera",
            "Nivel",
            "Hora de llegada"
        ])

        workbook.save(archivo)
        workbook.close()

        mensaje.config(
            text="Registro del día generado correctamente.",
            fg=COLOR_EXITO
        )

    # Activar lector
    entrada_matricula.config(
        state="normal"
    )

    entrada_matricula.focus()

    # Ocultar botón
    boton_generar.place_forget()


# ==========================================
# BUSCAR ALUMNO
# ==========================================

def buscar_alumno():

    # Si está bloqueado, no hacer nada
    if entrada_matricula["state"] == "disabled":
        return

    matricula = entrada_matricula.get().strip()

    # Limpiar inmediatamente
    entrada_matricula.delete(
        0,
        tk.END
    )

    if not matricula:

        entrada_matricula.focus()

        return

    # ======================================
    # ABRIR ARCHIVO DE ALUMNOS
    # ======================================

    workbook = load_workbook(
        ALUMNOS_FILE,
        data_only=True
    )

    hoja = workbook["Alumnos"]

    alumno_encontrado = None

    # ======================================
    # BUSCAR MATRÍCULA
    # ======================================

    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):

        matricula_excel = str(
            fila[0]
        ).strip()

        if matricula_excel == matricula:

            alumno_encontrado = fila

            break

    workbook.close()

    # ======================================
    # ALUMNO NO ENCONTRADO
    # ======================================

    if alumno_encontrado is None:

        mostrar_resultado(
            "ALUMNO NO ENCONTRADO",
            "La matrícula no está registrada.",
            COLOR_ERROR
        )

        return

    # ======================================
    # DATOS DEL ALUMNO
    # ======================================

    matricula_alumno, nombre, carrera, nivel = alumno_encontrado

    # ======================================
    # ARCHIVO DE ASISTENCIA
    # ======================================

    archivo_asistencia = obtener_archivo_asistencia()

    # Si no existe, crearlo automáticamente
    if not archivo_asistencia.exists():

        workbook = Workbook()

        hoja = workbook.active
        hoja.title = "Asistencias"

        hoja.append([
            "Matrícula",
            "Nombre",
            "Carrera",
            "Nivel",
            "Hora de llegada"
        ])

        workbook.save(
            archivo_asistencia
        )

        workbook.close()

    # ======================================
    # COMPROBAR SI YA ESTÁ REGISTRADO
    # ======================================

    workbook = load_workbook(
        archivo_asistencia
    )

    hoja = workbook["Asistencias"]

    ya_registrado = False
    hora_anterior = None

    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):

        matricula_registrada = str(
            fila[0]
        ).strip()

        if matricula_registrada == str(
            matricula
        ).strip():

            ya_registrado = True

            hora_anterior = fila[4]

            break

    # ======================================
    # YA REGISTRADO
    # ======================================

    if ya_registrado:

        workbook.close()

        mostrar_resultado(
            "YA REGISTRADO",
            f"{nombre}\n\nEntrada registrada: {hora_anterior}",
            COLOR_ADVERTENCIA
        )

        return

    # ======================================
    # REGISTRAR ASISTENCIA
    # ======================================

    hora_actual = datetime.now().strftime(
        "%H:%M:%S"
    )

    hoja.append([
        matricula_alumno,
        nombre,
        carrera,
        nivel,
        hora_actual
    ])

    workbook.save(
        archivo_asistencia
    )

    workbook.close()

    # ======================================
    # MOSTRAR BIENVENIDA
    # ======================================

    mostrar_bienvenida(
        nombre,
        hora_actual
    )


# ==========================================
# MOSTRAR BIENVENIDA
# ==========================================

def mostrar_bienvenida(
    nombre,
    hora
):

    # Ocultar controles
    entrada_matricula.place_forget()
    boton_generar.place_forget()

    # Título
    titulo.config(
        text="BIENVENIDO A",
        font=("Arial", 28, "bold"),
        fg=COLOR_TEXTO
    )

    # Institución
    subtitulo.config(
        text="INSUCO APODACA",
        font=("Arial", 34, "bold"),
        fg=COLOR_TEXTO
    )

    # Información
    mensaje.config(
        text=f"{nombre}\n\nEntrada registrada\n{hora}",
        font=("Arial", 18),
        fg=COLOR_EXITO
    )

    # Regresar después de 2 segundos
    ventana.after(
        2000,
        restaurar_pantalla
    )


# ==========================================
# MOSTRAR RESULTADO
# ==========================================

def mostrar_resultado(
    titulo_texto,
    mensaje_texto,
    color
):

    # Ocultar controles
    entrada_matricula.place_forget()
    boton_generar.place_forget()

    # --------------------------------------
    # TÍTULO
    # --------------------------------------

    titulo.config(
        text=titulo_texto,
        font=("Arial", 28, "bold"),
        fg=COLOR_TEXTO
    )

    titulo.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.35,
        anchor="center"
    )

    # --------------------------------------
    # NOMBRE / MENSAJE PRINCIPAL
    # --------------------------------------

    partes = mensaje_texto.split("\n\n", 1)

    nombre = partes[0]

    informacion = ""

    if len(partes) > 1:
        informacion = partes[1]

    subtitulo.config(
        text=nombre,
        font=("Arial", 22, "bold"),
        fg=color
    )

    subtitulo.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.46,
        anchor="center"
    )

    # --------------------------------------
    # INFORMACIÓN
    # --------------------------------------

    mensaje.config(
        text=informacion,
        font=("Arial", 17),
        fg=color
    )

    mensaje.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.56,
        anchor="center"
    )

    # Bloquear lector
    entrada_matricula.config(
        state="disabled"
    )

    ventana.after(
        2000,
        restaurar_pantalla
    )


# ==========================================
# RESTAURAR PANTALLA
# ==========================================

def restaurar_pantalla():

    # --------------------------------------
    # TÍTULO
    # --------------------------------------

    titulo.config(
        text="CONTROL DE ASISTENCIA",
        font=("Arial", 27, "bold"),
        fg=COLOR_TEXTO
    )

    titulo.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.36,
        anchor="center"
    )

    # --------------------------------------
    # SUBTÍTULO
    # --------------------------------------

    subtitulo.config(
        text="Escanea tu credencial",
        font=("Arial", 17),
        fg=COLOR_SECUNDARIO
    )

    subtitulo.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.44,
        anchor="center"
    )

    # --------------------------------------
    # INPUT
    # --------------------------------------

    entrada_matricula.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.57,
        anchor="center",
        width=int(CARD_WIDTH * 0.76),
        height=55
    )

    entrada_matricula.config(
        state="normal"
    )

    entrada_matricula.delete(
        0,
        tk.END
    )

    # --------------------------------------
    # MENSAJE
    # --------------------------------------

    mensaje.config(
        text="Esperando escaneo...",
        font=("Arial", 15),
        fg=COLOR_SECUNDARIO
    )

    mensaje.place(
        x=CARD_X + CARD_WIDTH * 0.50,
        y=CARD_Y + CARD_HEIGHT * 0.71,
        anchor="center"
    )

    entrada_matricula.focus()

# ==========================================
# RELOJ Y FONDO
# ==========================================

def obtener_fuente_reloj(tamano, negrita=False):
    """
    Busca una fuente disponible tanto en Linux como en Windows.
    """

    if negrita:
        posibles = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
            "C:/Windows/Fonts/arialbd.ttf",
            "C:/Windows/Fonts/segoeuib.ttf"
        ]
    else:
        posibles = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/segoeui.ttf"
        ]

    for ruta in posibles:
        if Path(ruta).exists():
            return ImageFont.truetype(ruta, tamano)

    return ImageFont.load_default()


def crear_imagen_fondo():
    """
    Construye el fondo completo en memoria:
    fotografía + oscurecimiento + capa azul +
    tarjeta + fecha + hora.

    La fecha y la hora se dibujan directamente
    sobre la imagen, por lo que NO tienen ningún
    Label ni rectángulo detrás.
    """

    imagen = Image.open(
        FONDO_FILE
    ).convert("RGB")

    # Ajustar fotografía a la resolución de pantalla
    imagen = ImageOps.fit(
        imagen,
        (ANCHO, ALTO),
        method=Image.Resampling.LANCZOS
    )

    # Oscurecer fotografía
    mejorador = ImageEnhance.Brightness(
        imagen
    )

    imagen = mejorador.enhance(
        0.50
    )

    # Capa azul
    capa_azul = Image.new(
        "RGB",
        (ANCHO, ALTO),
        "#07162e"
    )

    imagen = Image.blend(
        imagen,
        capa_azul,
        0.35
    )

    # Tarjeta principal
    dibujar = ImageDraw.Draw(
        imagen
    )

    dibujar.rounded_rectangle(
        (
            CARD_X,
            CARD_Y,
            CARD_X + CARD_WIDTH,
            CARD_Y + CARD_HEIGHT
        ),
        radius=RADIO,
        fill=COLOR_PANEL
    )

    # ======================================
    # FECHA Y HORA
    # ======================================

    ahora = datetime.now()

    dias = [
        "Lunes",
        "Martes",
        "Miércoles",
        "Jueves",
        "Viernes",
        "Sábado",
        "Domingo"
    ]

    meses = [
        "enero",
        "febrero",
        "marzo",
        "abril",
        "mayo",
        "junio",
        "julio",
        "agosto",
        "septiembre",
        "octubre",
        "noviembre",
        "diciembre"
    ]

    fecha = (
        f"{dias[ahora.weekday()]}, "
        f"{ahora.day} de "
        f"{meses[ahora.month - 1]} de "
        f"{ahora.year}"
    )

    hora = ahora.strftime("%H:%M")

    fuente_fecha = obtener_fuente_reloj(
        16,
        False
    )

    fuente_hora = obtener_fuente_reloj(
        40,
        True
    )

    # Margen desde el borde derecho
    margen_derecho = 48

    # ======================================
    # POSICIÓN FECHA
    # ======================================

    caja_fecha = dibujar.textbbox(
        (0, 0),
        fecha,
        font=fuente_fecha
    )

    ancho_fecha = (
        caja_fecha[2] - caja_fecha[0]
    )

    x_fecha = (
        ANCHO
        - margen_derecho
        - ancho_fecha
    )

    y_fecha = 30

    # ======================================
    # POSICIÓN HORA
    # ======================================

    caja_hora = dibujar.textbbox(
        (0, 0),
        hora,
        font=fuente_hora
    )

    ancho_hora = (
        caja_hora[2] - caja_hora[0]
    )

    x_hora = (
        ANCHO
        - margen_derecho
        - ancho_hora
    )

    y_hora = 55

    # ======================================
    # SOMBRA MUY SUTIL
    # ======================================

    dibujar.text(
        (
            x_fecha + 2,
            y_fecha + 2
        ),
        fecha,
        font=fuente_fecha,
        fill="#000000"
    )

    dibujar.text(
        (
            x_hora + 3,
            y_hora + 3
        ),
        hora,
        font=fuente_hora,
        fill="#000000"
    )

    # ======================================
    # TEXTO BLANCO
    # ======================================

    dibujar.text(
        (
            x_fecha,
            y_fecha
        ),
        fecha,
        font=fuente_fecha,
        fill="#ffffff"
    )

    dibujar.text(
        (
            x_hora,
            y_hora
        ),
        hora,
        font=fuente_hora,
        fill="#ffffff"
    )

    return imagen


def actualizar_reloj():
    """
    Regenera únicamente la composición visual del fondo
    cada segundo para mantener actualizada la hora.
    """

    global fondo_tk

    imagen_actual = crear_imagen_fondo()

    fondo_tk = ImageTk.PhotoImage(
        imagen_actual
    )

    fondo.config(
        image=fondo_tk
    )

    fondo.image = fondo_tk

    ventana.after(
        1000,
        actualizar_reloj
    )


# ==========================================
# CREAR ARCHIVO DE ALUMNOS
# ==========================================

crear_archivo_alumnos()


# ==========================================
# VENTANA PRINCIPAL
# ==========================================

ventana = tk.Tk()

ventana.title(
    "Control de Asistencia"
)

# Pantalla completa
ventana.attributes(
    "-fullscreen",
    True
)

# ESC para cerrar durante desarrollo
ventana.bind(
    "<Escape>",
    lambda evento: ventana.destroy()
)


# ==========================================
# RESOLUCIÓN
# ==========================================

ANCHO = ventana.winfo_screenwidth()
ALTO = ventana.winfo_screenheight()


# ==========================================
# DIMENSIONES DE TARJETA
# ==========================================

CARD_WIDTH = int(
    ANCHO * 0.43
)

CARD_HEIGHT = int(
    ALTO * 0.78
)

CARD_X = int(
    (ANCHO - CARD_WIDTH) / 2
)

CARD_Y = int(
    (ALTO - CARD_HEIGHT) / 2
)

RADIO = 30


# ==========================================
# FONDO INICIAL
# ==========================================

imagen_fondo = crear_imagen_fondo()

fondo_tk = ImageTk.PhotoImage(
    imagen_fondo
)

fondo = tk.Label(
    ventana,
    image=fondo_tk,
    borderwidth=0,
    highlightthickness=0
)

fondo.place(
    x=0,
    y=0,
    relwidth=1,
    relheight=1
)


# ==========================================
# LOGO
# ==========================================

imagen_logo = Image.open(
    LOGO_FILE
).convert("RGBA")

imagen_logo = imagen_logo.resize(
    (247, 103),
    Image.Resampling.LANCZOS
)

logo_tk = ImageTk.PhotoImage(
    imagen_logo
)

logo = tk.Label(
    ventana,
    image=logo_tk,
    bg=COLOR_PANEL,
    borderwidth=0,
    highlightthickness=0
)

logo.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.14,
    anchor="center"
)


# ==========================================
# LÍNEA DECORATIVA
# ==========================================

linea = tk.Frame(
    ventana,
    bg=COLOR_AZUL,
    width=80,
    height=2,
    bd=0
)

linea.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.275,
    anchor="center"
)


# ==========================================
# TÍTULO
# ==========================================

titulo = tk.Label(
    ventana,
    text="CONTROL DE ASISTENCIA",
    font=("Arial", 27, "bold"),
    fg=COLOR_TEXTO,
    bg=COLOR_PANEL,
    borderwidth=0,
    highlightthickness=0
)

titulo.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.36,
    anchor="center"
)


# ==========================================
# SUBTÍTULO
# ==========================================

subtitulo = tk.Label(
    ventana,
    text="Escanea tu credencial",
    font=("Arial", 17),
    fg=COLOR_SECUNDARIO,
    bg=COLOR_PANEL,
    borderwidth=0,
    highlightthickness=0
)

subtitulo.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.44,
    anchor="center"
)


# ==========================================
# INPUT QR
# ==========================================

entrada_matricula = tk.Entry(
    ventana,
    font=("Arial", 22),
    fg=COLOR_TEXTO,
    bg=COLOR_INPUT,
    insertbackground=COLOR_TEXTO,
    justify="center",
    relief="flat",
    bd=0,
    highlightbackground=COLOR_AZUL,
    highlightcolor=COLOR_AZUL,
    highlightthickness=2
)

entrada_matricula.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.57,
    anchor="center",
    width=int(CARD_WIDTH * 0.76),
    height=55
)


# ==========================================
# BOTÓN GENERAR
# ==========================================

boton_generar = tk.Button(
    ventana,
    text="GENERAR REGISTRO DE HOY",
    font=("Arial", 14, "bold"),
    fg=COLOR_TEXTO,
    bg="#2868c7",
    activebackground="#347bdc",
    activeforeground=COLOR_TEXTO,
    relief="flat",
    bd=0,
    cursor="hand2",
    command=generar_archivo_dia
)

boton_generar.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.68,
    anchor="center",
    width=int(CARD_WIDTH * 0.76),
    height=58
)


# ==========================================
# MENSAJE
# ==========================================

mensaje = tk.Label(
    ventana,
    text="",
    font=("Arial", 15),
    fg=COLOR_SECUNDARIO,
    bg=COLOR_PANEL,
    borderwidth=0,
    highlightthickness=0
)

mensaje.place(
    x=CARD_X + CARD_WIDTH * 0.50,
    y=CARD_Y + CARD_HEIGHT * 0.76,
    anchor="center"
)


# ==========================================
# ESTADO INICIAL
# ==========================================

archivo_hoy = obtener_archivo_asistencia()

if archivo_hoy.exists():

    # Ya existe el registro de hoy.
    # Ocultar botón.

    boton_generar.place_forget()

    entrada_matricula.config(
        state="normal"
    )

    mensaje.config(
        text="Esperando escaneo...",
        fg=COLOR_SECUNDARIO
    )

else:

    # Todavía no existe registro.

    entrada_matricula.config(
        state="disabled"
    )

    mensaje.config(
        text="Presiona el botón para comenzar.",
        fg=COLOR_SECUNDARIO
    )


# ==========================================
# ENTER DEL LECTOR QR
# ==========================================

ventana.bind(
    "<Return>",
    lambda evento: buscar_alumno()
)


# ==========================================
# INICIAR RELOJ
# ==========================================

actualizar_reloj()


# ==========================================
# ENFOCAR INPUT
# ==========================================

if archivo_hoy.exists():

    entrada_matricula.focus()


# ==========================================
# INICIO
# ==========================================

ventana.mainloop()
